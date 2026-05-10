"""
文档加载模块
支持 Markdown、PDF、Word、TXT 格式
"""

from typing import List, Dict, Any, Optional
from pathlib import Path
import yaml


class DocumentLoader:
    """文档加载器基类"""

    def load(self, file_path: str) -> Dict[str, Any]:
        """
        加载单个文档

        Args:
            file_path: 文件路径

        Returns:
            文档字典，包含 content 和 metadata
        """
        raise NotImplementedError

    def load_batch(self, file_paths: List[str]) -> List[Dict[str, Any]]:
        """
        批量加载文档

        Args:
            file_paths: 文件路径列表

        Returns:
            文档列表
        """
        documents = []
        for file_path in file_paths:
            try:
                doc = self.load(file_path)
                documents.append(doc)
            except Exception as e:
                print(f"❌ 加载失败: {file_path}, 错误: {e}")
        return documents


class MarkdownLoader(DocumentLoader):
    """Markdown 文档加载器"""

    def load(self, file_path: str) -> Dict[str, Any]:
        """加载 Markdown 文件"""
        path = Path(file_path)

        with open(path, 'r', encoding='utf-8') as f:
            content = f.read()

        # 提取 frontmatter（YAML 格式）
        metadata = self._extract_frontmatter(content)

        # 移除 frontmatter
        content = self._remove_frontmatter(content)

        # 添加文件信息
        metadata.update({
            "source": str(path),
            "file_name": path.name,
            "file_type": "markdown"
        })

        return {
            "content": content.strip(),
            "metadata": metadata
        }

    def _extract_frontmatter(self, content: str) -> Dict[str, Any]:
        """提取 frontmatter"""
        if not content.startswith("---"):
            return {}

        try:
            # 查找第二个 ---
            end_index = content.find("---", 3)
            if end_index == -1:
                return {}

            frontmatter_text = content[3:end_index].strip()
            metadata = yaml.safe_load(frontmatter_text)

            # 如果有嵌套的 metadata 字段，提取出来
            if isinstance(metadata, dict) and "metadata" in metadata:
                return metadata["metadata"]

            return metadata if isinstance(metadata, dict) else {}

        except Exception as e:
            print(f"⚠️  解析 frontmatter 失败: {e}")
            return {}

    def _remove_frontmatter(self, content: str) -> str:
        """移除 frontmatter"""
        if not content.startswith("---"):
            return content

        end_index = content.find("---", 3)
        if end_index == -1:
            return content

        return content[end_index + 3:].strip()


class PDFLoader(DocumentLoader):
    """PDF 文档加载器"""

    def load(self, file_path: str) -> Dict[str, Any]:
        """加载 PDF 文件"""
        from pypdf import PdfReader

        path = Path(file_path)
        reader = PdfReader(path)

        # 提取所有页面文本
        content = ""
        for page in reader.pages:
            content += page.extract_text() + "\n\n"

        # 提取元数据
        metadata = {
            "source": str(path),
            "file_name": path.name,
            "file_type": "pdf",
            "total_pages": len(reader.pages)
        }

        # 尝试从 PDF 元数据中提取信息
        if reader.metadata:
            pdf_meta = reader.metadata
            if pdf_meta.title:
                metadata["title"] = pdf_meta.title
            if pdf_meta.author:
                metadata["author"] = pdf_meta.author

        return {
            "content": content.strip(),
            "metadata": metadata
        }


class WordLoader(DocumentLoader):
    """Word 文档加载器"""

    def load(self, file_path: str) -> Dict[str, Any]:
        """加载 Word 文件"""
        from docx import Document

        path = Path(file_path)
        doc = Document(path)

        # 提取所有段落文本
        content = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])

        metadata = {
            "source": str(path),
            "file_name": path.name,
            "file_type": "docx",
            "total_paragraphs": len(doc.paragraphs)
        }

        return {
            "content": content.strip(),
            "metadata": metadata
        }


class TextLoader(DocumentLoader):
    """纯文本加载器"""

    def load(self, file_path: str) -> Dict[str, Any]:
        """加载纯文本文件"""
        path = Path(file_path)

        with open(path, 'r', encoding='utf-8') as f:
            content = f.read()

        metadata = {
            "source": str(path),
            "file_name": path.name,
            "file_type": "txt"
        }

        return {
            "content": content.strip(),
            "metadata": metadata
        }


def get_loader(file_path: str) -> DocumentLoader:
    """
    根据文件扩展名获取对应的加载器

    Args:
        file_path: 文件路径

    Returns:
        文档加载器实例
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    loaders = {
        ".md": MarkdownLoader,
        ".markdown": MarkdownLoader,
        ".pdf": PDFLoader,
        ".docx": WordLoader,
        ".doc": WordLoader,
        ".txt": TextLoader
    }

    loader_class = loaders.get(ext)
    if not loader_class:
        raise ValueError(f"不支持的文件格式: {ext}")

    return loader_class()


def load_document(file_path: str) -> Dict[str, Any]:
    """
    便捷函数：加载单个文档

    Args:
        file_path: 文件路径

    Returns:
        文档字典
    """
    loader = get_loader(file_path)
    return loader.load(file_path)


def load_documents(file_paths: List[str]) -> List[Dict[str, Any]]:
    """
    便捷函数：批量加载文档

    Args:
        file_paths: 文件路径列表

    Returns:
        文档列表
    """
    documents = []
    for file_path in file_paths:
        try:
            doc = load_document(file_path)
            documents.append(doc)
            print(f"✅ 加载成功: {file_path}")
        except Exception as e:
            print(f"❌ 加载失败: {file_path}, 错误: {e}")

    return documents


if __name__ == "__main__":
    # 测试
    import sys

    if len(sys.argv) > 1:
        file_path = sys.argv[1]
        print(f"测试加载文档: {file_path}")

        doc = load_document(file_path)
        print(f"\n✅ 加载成功")
        print(f"   内容长度: {len(doc['content'])} 字符")
        print(f"   元数据: {doc['metadata']}")
        print(f"\n内容预览:")
        print(doc['content'][:500])
    else:
        print("用法: python loaders.py <file_path>")
